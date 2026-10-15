import io

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .dashboard import MESES

CABECALHO_NOTIF = ["Filial", "Nº DET", "Assunto", "Recebimento", "Prazo", "Status"]


def _linha_notificacao(n):
    return [
        n.filial.nome,
        n.numero_det,
        n.assunto,
        n.data_recebimento.strftime("%d/%m/%Y"),
        n.prazo_resposta.strftime("%d/%m/%Y") if n.prazo_resposta else "",
        n.status_exibicao,
    ]


def _xlsx_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _escrever(ws, cabecalho, linhas):
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append(linha)
    for coluna in ws.columns:
        ws.column_dimensions[coluna[0].column_letter].width = min(
            max(len(str(c.value or "")) for c in coluna) + 2, 60
        )


def _pdf_bytes(titulo, secoes, paisagem=True):
    buf = io.BytesIO()
    pagina = landscape(A4) if paisagem else A4
    doc = SimpleDocTemplate(buf, pagesize=pagina, title=titulo, leftMargin=24, rightMargin=24)
    estilos = getSampleStyleSheet()
    elementos = []
    logo = settings.BASE_DIR / "static" / "img" / "logo_isac.png"
    if logo.exists():
        elementos.append(Image(str(logo), width=150, height=150 * 790 / 2000, hAlign="LEFT"))
    elementos += [Paragraph(titulo, estilos["Title"]), Spacer(1, 8)]
    for subtitulo, cabecalho, linhas in secoes:
        if subtitulo:
            elementos += [Paragraph(subtitulo, estilos["Heading3"])]
        dados = [cabecalho] + [[Paragraph(str(c), estilos["BodyText"]) for c in l] for l in linhas]
        tabela = Table(dados, repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#005aa9")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elementos += [tabela, Spacer(1, 12)]
    doc.build(elementos)
    return buf.getvalue()


def notificacoes_xlsx(qs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Notificações"
    _escrever(ws, CABECALHO_NOTIF, [_linha_notificacao(n) for n in qs])
    return _xlsx_bytes(wb)


def notificacoes_pdf(qs):
    return _pdf_bytes("Notificações do DET", [(None, CABECALHO_NOTIF, [_linha_notificacao(n) for n in qs])])


def _linhas_filiais(dados):
    return [
        [f["filial"].nome, f["total_ano"], f["total_mes"], f["vencidas"], f["situacao_rotulo"]]
        for f in dados["filiais"]
    ]


CABECALHO_FILIAIS = ["Filial", "Total no ano", "Total no mês", "Vencidas", "Situação no mês"]


def dashboard_xlsx(dados):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    r = dados["resumo"]
    _escrever(ws, ["Indicador", "Valor"], [
        ["Competência", f"{dados['mes']:02d}/{dados['ano']}"],
        ["Total no ano", r["total_ano"]], ["Total no mês", r["total_mes"]],
        ["Pendentes", r["pendentes"]], ["Em andamento", r["em_andamento"]],
        ["Respondidas", r["respondidas"]], ["Vencidas", r["vencidas"]],
        ["Prazo nos próximos 7 dias", r["prazo_7_dias"]],
    ])
    _escrever(wb.create_sheet("Por filial"), CABECALHO_FILIAIS, _linhas_filiais(dados))
    _escrever(wb.create_sheet("Série mensal"), ["Mês", "Total"], [[s["mes"], s["total"]] for s in dados["serie"]])
    return _xlsx_bytes(wb)


def dashboard_pdf(dados):
    r = dados["resumo"]
    resumo = [
        ["Total no ano", r["total_ano"]], ["Total no mês", r["total_mes"]],
        ["Pendentes", r["pendentes"]], ["Em andamento", r["em_andamento"]],
        ["Respondidas", r["respondidas"]], ["Vencidas", r["vencidas"]],
    ]
    return _pdf_bytes(
        f"Dashboard DET - {MESES[dados['mes'] - 1]}/{dados['ano']}",
        [
            ("Resumo", ["Indicador", "Valor"], resumo),
            ("Por filial", CABECALHO_FILIAIS, _linhas_filiais(dados)),
            ("Série mensal", ["Mês", "Total"], [[s["mes"], s["total"]] for s in dados["serie"]]),
        ],
    )
