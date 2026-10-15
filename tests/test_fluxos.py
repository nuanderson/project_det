import datetime
import io

import pytest
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.urls import reverse
from openpyxl import load_workbook

from core.models import LogAuditoria, Notificacao

pytestmark = pytest.mark.django_db


def _dados_form(**extra):
    base = {
        "numero_det": "DET-1", "assunto": "Solicitação", "data_recebimento": "2025-05-10",
        "prazo_resposta": "2025-05-20", "status": "PENDENTE", "observacao": "",
    }
    base.update(extra)
    return base


def test_filial_cadastra_notificacao_com_anexo(client, user_a, filial_a):
    client.force_login(user_a)
    pdf = SimpleUploadedFile("oficio.pdf", b"%PDF-1.4 teste", content_type="application/pdf")
    resp = client.post(reverse("notificacao_criar"), _dados_form(anexo=pdf))
    assert resp.status_code == 302
    n = Notificacao.objects.get()
    assert n.filial == filial_a and n.criado_por == user_a and n.anexo
    assert LogAuditoria.objects.filter(acao="criar_notificacao").exists()
    baixar = client.get(reverse("notificacao_anexo", args=[n.pk]))
    assert baixar.status_code == 200 and b"".join(baixar.streaming_content).startswith(b"%PDF")


def test_filial_nao_consegue_escolher_outra_filial(client, user_a, filial_a, filial_b):
    client.force_login(user_a)
    client.post(reverse("notificacao_criar"), _dados_form(filial=filial_b.pk))
    assert Notificacao.objects.get().filial == filial_a


def test_anexo_com_extensao_proibida_e_rejeitado(client, user_a):
    client.force_login(user_a)
    exe = SimpleUploadedFile("virus.exe", b"MZ", content_type="application/octet-stream")
    resp = client.post(reverse("notificacao_criar"), _dados_form(anexo=exe))
    assert resp.status_code == 200 and not Notificacao.objects.exists()


def test_anexo_acima_do_limite_e_rejeitado(client, user_a, settings):
    settings.ANEXO_MAX_MB = 0  # qualquer arquivo não vazio excede
    client.force_login(user_a)
    grande = SimpleUploadedFile("a.pdf", b"x" * 10, content_type="application/pdf")
    resp = client.post(reverse("notificacao_criar"), _dados_form(anexo=grande))
    assert resp.status_code == 200 and not Notificacao.objects.exists()


def test_prazo_anterior_ao_recebimento_e_rejeitado(client, user_a):
    client.force_login(user_a)
    resp = client.post(reverse("notificacao_criar"), _dados_form(prazo_resposta="2025-05-01"))
    assert resp.status_code == 200 and not Notificacao.objects.exists()


def test_dashboard_e_exportacoes_da_matriz(client, matriz, user_a, filial_a, criar_notificacao):
    criar_notificacao(filial_a, user_a, data=datetime.date(2025, 6, 5), numero="Z1")
    client.force_login(matriz)
    resp = client.get(reverse("dashboard"), {"ano": 2025, "mes": 6})
    assert resp.status_code == 200 and b"Filial A" in resp.content

    xlsx = client.get(reverse("exportar_dashboard", args=["xlsx"]), {"ano": 2025, "mes": 6})
    assert "Por filial" in load_workbook(io.BytesIO(xlsx.content)).sheetnames
    pdf = client.get(reverse("exportar_dashboard", args=["pdf"]), {"ano": 2025, "mes": 6})
    assert pdf.content.startswith(b"%PDF")
    pdf_lista = client.get(reverse("exportar_notificacoes", args=["pdf"]))
    assert pdf_lista.content.startswith(b"%PDF")
    assert client.get(reverse("exportar_dashboard", args=["csv"])).status_code == 404


def test_periodo_invalido_cai_no_mes_atual(client, matriz):
    client.force_login(matriz)
    assert client.get(reverse("dashboard"), {"ano": "abc", "mes": "99"}).status_code == 200


def test_matriz_cria_usuario_de_filial_exigindo_filial(client, matriz, filial_a):
    client.force_login(matriz)
    dados = {"username": "novo", "perfil": "FILIAL", "password1": "Senha-Forte-987", "password2": "Senha-Forte-987"}
    assert client.post(reverse("usuario_criar"), dados).status_code == 200  # sem filial: erro
    assert client.post(reverse("usuario_criar"), {**dados, "filial": filial_a.pk}).status_code == 302


def test_superusuario_nasce_como_matriz(client):
    from core.models import Usuario
    admin = Usuario.objects.create_superuser("admin", "a@x.local", "Senha-Forte-987")
    assert admin.is_matriz and admin.filial is None
    client.force_login(admin)
    assert client.get(reverse("dashboard")).status_code == 200


def test_popular_demo_recusa_rodar_fora_do_debug(settings):
    from django.core.management.base import CommandError
    settings.DEBUG = False
    with pytest.raises(CommandError):
        call_command("popular_demo")


# ---------- alertas ----------

def test_alerta_lembrete_mensal_para_filiais_pendentes(matriz, user_a, filial_a, filial_b, criar_notificacao, settings):
    matriz.email = "matriz@x.local"
    matriz.save()
    user_a.email = "a@x.local"
    user_a.save()
    filial_a.email = "rh-a@x.local"
    filial_a.save()
    call_command("enviar_alertas", "--forcar-lembrete")
    destinos = {tuple(m.to) for m in mail.outbox}
    assert ("a@x.local", "rh-a@x.local") in destinos
    assert ("matriz@x.local",) in destinos
    resumo_matriz = next(m for m in mail.outbox if m.to == ["matriz@x.local"])
    assert "Filial A" in resumo_matriz.body and "Filial B" in resumo_matriz.body


def test_alerta_de_prazo_e_vencidas(matriz, user_a, filial_a, criar_notificacao, hoje, settings):
    settings.ALERTA_DIAS_ANTES_PRAZO = [2]
    matriz.email = "matriz@x.local"
    matriz.save()
    user_a.email = "a@x.local"
    user_a.save()
    criar_notificacao(filial_a, user_a, prazo=hoje + datetime.timedelta(days=2), numero="P2")
    criar_notificacao(filial_a, user_a, prazo=hoje - datetime.timedelta(days=3), numero="V1")
    call_command("enviar_alertas")
    assuntos = [m.subject for m in mail.outbox]
    assert any("prazo em 2 dia" in a for a in assuntos)
    assert any("vencidas" in a for a in assuntos)
    prazo = next(m for m in mail.outbox if "prazo" in m.subject)
    assert "P2" in prazo.body and "V1" not in prazo.body
