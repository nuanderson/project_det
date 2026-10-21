import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from core.forms import cnpj_valido
from core.models import Filial, Notificacao

pytestmark = pytest.mark.django_db


def _lista(client, **params):
    return client.get(reverse("notificacao_lista"), params)


def _numeros(resp):
    return [n.numero_det for n in resp.context["notificacoes"]]


# ---------- busca ----------

def test_busca_por_numero_assunto_processo_e_filial(client, matriz, user_a, user_b, filial_a, filial_b, criar_notificacao):
    criar_notificacao(filial_a, user_a, numero="ALFA-1")
    n2 = criar_notificacao(filial_a, user_a, numero="BETA-2")
    n2.assunto, n2.numero_processo = "Folha de pagamento", "12345678901202617"
    n2.save()
    criar_notificacao(filial_b, user_b, numero="GAMA-3")
    client.force_login(matriz)
    assert _numeros(_lista(client, q="alfa")) == ["ALFA-1"]
    assert _numeros(_lista(client, q="folha")) == ["BETA-2"]
    assert _numeros(_lista(client, q="12345678901")) == ["BETA-2"]
    assert _numeros(_lista(client, q="Filial B")) == ["GAMA-3"]
    assert "Nenhuma notificação neste período ou filtro" in _lista(client, q="inexistente").content.decode()


def test_busca_da_filial_nao_enxerga_outras(client, user_a, user_b, filial_a, filial_b, criar_notificacao):
    criar_notificacao(filial_a, user_a, numero="X-A")
    criar_notificacao(filial_b, user_b, numero="X-B")
    client.force_login(user_a)
    assert _numeros(_lista(client, q="X-")) == ["X-A"]


# ---------- ordenação ----------

def test_ordenacao_por_prazo_poe_sem_prazo_no_fim(client, matriz, user_a, filial_a, criar_notificacao, hoje, dias):
    criar_notificacao(filial_a, user_a, prazo=hoje + dias(9), numero="LONGE")
    criar_notificacao(filial_a, user_a, prazo=None, numero="SEM")
    criar_notificacao(filial_a, user_a, prazo=hoje + dias(1), numero="PERTO")
    client.force_login(matriz)
    assert _numeros(_lista(client, ordem="prazo")) == ["PERTO", "LONGE", "SEM"]
    assert _numeros(_lista(client, ordem="-prazo")) == ["LONGE", "PERTO", "SEM"]


def test_ordenacao_padrao_e_valor_invalido_nao_quebram(client, matriz, user_a, filial_a, criar_notificacao, hoje, dias):
    criar_notificacao(filial_a, user_a, data=hoje - dias(5), numero="VELHA")
    criar_notificacao(filial_a, user_a, data=hoje, numero="NOVA")
    client.force_login(matriz)
    assert _numeros(_lista(client)) == ["NOVA", "VELHA"]
    resp = _lista(client, ordem="numero_det; DROP TABLE")
    assert resp.status_code == 200 and _numeros(resp) == ["NOVA", "VELHA"]
    assert _numeros(_lista(client, ordem="numero")) == ["NOVA", "VELHA"]  # alfabética: NOVA < VELHA
    assert _numeros(_lista(client, ordem="-numero")) == ["VELHA", "NOVA"]


def test_cabecalho_indica_a_coluna_ordenada_e_inverte_ao_clicar(client, matriz):
    client.force_login(matriz)
    html = _lista(client, ordem="prazo").content.decode()
    assert 'aria-sort="ascending"' in html and "ordem=-prazo" in html
    html = _lista(client, ordem="-prazo").content.decode()
    assert 'aria-sort="descending"' in html and "ordem=prazo" in html


# ---------- paginação ----------

def test_paginacao_no_servidor(client, matriz, user_a, filial_a, criar_notificacao, settings):
    settings.PAGE_SIZE = 10
    for i in range(25):
        criar_notificacao(filial_a, user_a, numero=f"N{i:02d}")
    client.force_login(matriz)
    p1 = _lista(client, q="N")
    assert len(_numeros(p1)) == 10 and "Mostrando 1–10 de 25" in p1.content.decode()
    p3 = _lista(client, q="N", page=3)
    assert len(_numeros(p3)) == 5 and "Mostrando 21–25 de 25" in p3.content.decode()
    assert len(_numeros(_lista(client, q="N", page=999))) == 5     # fora do intervalo: última página
    assert len(_numeros(_lista(client, q="N", page="abc"))) == 10  # inválida: primeira página


def test_links_de_pagina_mantem_busca_e_ordem(client, matriz, user_a, filial_a, criar_notificacao, settings):
    settings.PAGE_SIZE = 5
    for i in range(12):
        criar_notificacao(filial_a, user_a, numero=f"K{i:02d}")
    client.force_login(matriz)
    html = _lista(client, q="K", ordem="numero").content.decode()
    assert "page=2" in html and "q=K" in html and "ordem=numero" in html


def test_exportacao_leva_todos_os_resultados_sem_paginar(client, matriz, user_a, filial_a, criar_notificacao, settings):
    settings.PAGE_SIZE = 5
    for i in range(12):
        criar_notificacao(filial_a, user_a, numero=f"E{i:02d}")
    client.force_login(matriz)
    html = _lista(client, q="E", page=2).content.decode()
    link = html.split("exportar/xlsx/")[1].split('"')[0]
    assert link.startswith("?q=E") and "page=" not in link
    resp = client.get(reverse("exportar_notificacoes", args=["xlsx"]), {"q": "E"})
    assert len(list(load_workbook(io.BytesIO(resp.content)).active.iter_rows(min_row=2))) == 12


# ---------- destaque de prazo ----------

def test_destaque_de_prazo_vencido_proximo_e_hoje(client, matriz, user_a, filial_a, criar_notificacao, hoje, dias):
    criar_notificacao(filial_a, user_a, prazo=hoje - dias(3), numero="V")
    criar_notificacao(filial_a, user_a, prazo=hoje + dias(2), numero="P")
    criar_notificacao(filial_a, user_a, prazo=hoje, numero="H")
    criar_notificacao(filial_a, user_a, prazo=hoje - dias(9), status="RESPONDIDA", numero="R")
    client.force_login(matriz)
    html = _lista(client).content.decode()
    assert "vencido há 3 dias" in html and "vence em 2 dias" in html and "vence hoje" in html
    assert html.count('class="linha-vencida"') == 1  # a respondida não é destacada


def test_propriedades_de_prazo(user_a, filial_a, criar_notificacao, hoje, dias):
    n = criar_notificacao(filial_a, user_a, prazo=hoje - dias(1))
    assert n.dias_atraso == 1 and n.vencida and not n.prazo_proximo
    n.prazo_resposta = hoje + dias(7)
    assert n.prazo_proximo and n.dias_atraso == 0
    n.prazo_resposta = hoje + dias(8)
    assert not n.prazo_proximo
    n.prazo_resposta = None
    assert n.dias_para_prazo is None


# ---------- validações e máscaras ----------

def test_validador_de_cnpj():
    assert cnpj_valido("11222333000181") and cnpj_valido("11444777000161")
    assert not cnpj_valido("11222333000182") and not cnpj_valido("00000000000000") and not cnpj_valido("123")


def test_cadastro_de_filial_valida_cnpj_formata_e_evita_duplicado(client, matriz):
    client.force_login(matriz)
    dados = {"nome": "Filial Nova", "cnpj": "11222333000181", "cidade": "Teresina", "uf": "pi",
             "email": "rh@filial.org.br", "ativa": "on"}
    assert client.post(reverse("filial_nova"), {**dados, "cnpj": "11222333000182"}).status_code == 200  # dígito errado
    assert not Filial.objects.exists()
    assert client.post(reverse("filial_nova"), dados).status_code == 302
    f = Filial.objects.get()
    assert f.cnpj == "11.222.333/0001-81" and f.uf == "PI"
    duplicada = client.post(reverse("filial_nova"), {**dados, "nome": "Outra"})
    assert duplicada.status_code == 200 and Filial.objects.count() == 1
    invalido = client.post(reverse("filial_nova"), {**dados, "nome": "Outra", "cnpj": "11444777000161", "email": "invalido"})
    assert invalido.status_code == 200 and Filial.objects.count() == 1


def test_editar_filial_mantem_o_proprio_cnpj(client, matriz, filial_a):
    filial_a.cnpj = "11.222.333/0001-81"
    filial_a.save()
    client.force_login(matriz)
    dados = {"nome": "Filial A", "cnpj": "11.222.333/0001-81", "cidade": "X", "uf": "SP", "email": "", "ativa": "on"}
    assert client.post(reverse("filial_editar", args=[filial_a.pk]), dados).status_code == 302


def test_data_de_recebimento_no_futuro_e_recusada(client, user_a, hoje, dias):
    client.force_login(user_a)
    base = {"numero_det": "F1", "assunto": "x", "status": "PENDENTE"}
    client.post(reverse("notificacao_criar"), {**base, "data_recebimento": (hoje + dias(1)).isoformat()})
    assert not Notificacao.objects.exists()
    client.post(reverse("notificacao_criar"), {**base, "data_recebimento": hoje.isoformat()})
    assert Notificacao.objects.count() == 1


def test_campos_do_formulario_trazem_mascaras(client, user_a):
    client.force_login(user_a)
    html = client.get(reverse("notificacao_criar")).content.decode()
    assert 'data-mask="cnpj"' in html and 'data-mask="digitos"' in html and 'data-mask="codigo"' in html


# ---------- toasts e confirmação ----------

def test_mensagens_aparecem_como_toast(client, user_a):
    client.force_login(user_a)
    dados = {"numero_det": "T1", "assunto": "x", "status": "PENDENTE", "data_recebimento": "2026-01-10"}
    html = client.post(reverse("notificacao_criar"), dados, follow=True).content.decode()
    assert 'class="toast show' in html and "Notificação registrada." in html and 'data-autohide="true"' in html


def test_toast_de_erro_nao_some_sozinho(client, user_a, filial_a, hoje):
    client.force_login(user_a)
    resp = client.post(reverse("justificar", args=[filial_a.pk, hoje.year, hoje.month]), {"texto": "curta"}, follow=True)
    html = resp.content.decode()
    assert 'role="alert"' in html and 'data-autohide="false"' in html


def test_acoes_destrutivas_pedem_confirmacao(client, user_a, filial_a, criar_notificacao):
    n = criar_notificacao(filial_a, user_a)
    client.force_login(user_a)
    html = client.get(reverse("notificacao_editar", args=[n.pk])).content.decode()
    assert "data-confirmar=" in html and 'id="modal-confirmar"' in html
    client.post(reverse("declarar_sem", args=[filial_a.pk, 2020, 1]))
    assert "data-confirmar=" in client.get(reverse("painel"), {"ano": 2020, "mes": 1}).content.decode()


def test_css_e_js_levam_versao_para_evitar_cache_velho(client):
    html = client.get(reverse("login")).content.decode()
    assert "css/tema.css?v=" in html and "js/app.js?v=" in html
