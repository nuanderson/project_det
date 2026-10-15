import datetime

import pytest
from django.core.exceptions import PermissionDenied, ValidationError
from django.urls import reverse

from core.models import FechamentoMensal, JustificativaHistorico, Notificacao
from core.services import dashboard as dash
from core.services import fechamento as fech

pytestmark = pytest.mark.django_db


# ---------- isolamento entre filiais ----------

def test_filial_so_ve_as_proprias_notificacoes(client, user_a, user_b, filial_a, filial_b, criar_notificacao):
    na = criar_notificacao(filial_a, user_a, numero="A1")
    nb = criar_notificacao(filial_b, user_b, numero="B1")
    client.force_login(user_a)

    html = client.get(reverse("notificacao_lista")).content.decode()
    assert "A1" in html and "B1" not in html
    assert client.get(reverse("notificacao_editar", args=[nb.pk])).status_code == 404
    assert client.post(reverse("notificacao_excluir", args=[nb.pk])).status_code == 404
    assert client.get(reverse("notificacao_anexo", args=[nb.pk])).status_code == 404
    assert client.get(reverse("notificacao_editar", args=[na.pk])).status_code == 200


def test_filtro_por_filial_nao_vaza_para_usuario_de_filial(client, user_a, user_b, filial_a, filial_b, criar_notificacao):
    criar_notificacao(filial_b, user_b, numero="B1")
    client.force_login(user_a)
    html = client.get(reverse("notificacao_lista"), {"filial": filial_b.pk}).content.decode()
    assert "B1" not in html


def test_exportacao_respeita_escopo(client, user_a, user_b, filial_a, filial_b, criar_notificacao):
    criar_notificacao(filial_a, user_a, numero="A1")
    criar_notificacao(filial_b, user_b, numero="B1")
    client.force_login(user_a)
    from openpyxl import load_workbook
    import io
    resp = client.get(reverse("exportar_notificacoes", args=["xlsx"]))
    ws = load_workbook(io.BytesIO(resp.content)).active
    numeros = [row[1].value for row in ws.iter_rows(min_row=2)]
    assert numeros == ["A1"]


def test_filial_nao_acessa_telas_da_matriz(client, user_a):
    client.force_login(user_a)
    for nome in ("dashboard", "nao_enviaram", "filial_lista", "usuario_lista"):
        assert client.get(reverse(nome)).status_code == 302  # redireciona para login


def test_filial_nao_justifica_por_outra_filial(client, user_a, filial_b, hoje):
    client.force_login(user_a)
    with pytest.raises(PermissionDenied):
        fech.justificar(user_a, filial_b, hoje.year, hoje.month, "Justificativa longa o bastante")
    resp = client.post(reverse("declarar_sem", args=[filial_b.pk, hoje.year, hoje.month]))
    assert resp.status_code == 403


def test_matriz_ve_tudo(client, matriz, user_a, user_b, filial_a, filial_b, criar_notificacao):
    criar_notificacao(filial_a, user_a, numero="A1")
    criar_notificacao(filial_b, user_b, numero="B1")
    client.force_login(matriz)
    html = client.get(reverse("notificacao_lista")).content.decode()
    assert "A1" in html and "B1" in html


def test_next_externo_e_ignorado(client, user_a, filial_a, hoje):
    client.force_login(user_a)
    resp = client.post(
        reverse("declarar_sem", args=[filial_a.pk, hoje.year, hoje.month]), {"next": "https://evil.example/"}
    )
    assert resp.status_code == 302 and "evil.example" not in resp["Location"]


# ---------- fechamento mensal ----------

def test_transicoes_do_fechamento(user_a, filial_a, criar_notificacao, hoje):
    ano, mes = hoje.year, hoje.month
    f = fech.situacao_do_mes(ano, mes)[0]
    assert f["situacao"] == "PENDENTE"

    fech.justificar(user_a, filial_a, ano, mes, "Responsável em licença médica")
    assert FechamentoMensal.objects.get(filial=filial_a, ano=ano, mes=mes).situacao == "JUSTIFICADO"

    fech.declarar_sem_notificacoes(user_a, filial_a, ano, mes)
    assert FechamentoMensal.objects.get(filial=filial_a, ano=ano, mes=mes).situacao == "SEM_NOTIFICACOES"

    notif = criar_notificacao(filial_a, user_a, data=hoje)
    assert FechamentoMensal.objects.get(filial=filial_a, ano=ano, mes=mes).situacao == "COM_NOTIFICACOES"

    notif.delete()
    # Declaração continua valendo após excluir a notificação.
    assert FechamentoMensal.objects.get(filial=filial_a, ano=ano, mes=mes).situacao == "SEM_NOTIFICACOES"

    fech.desfazer_declaracao(user_a, filial_a, ano, mes)
    assert FechamentoMensal.objects.get(filial=filial_a, ano=ano, mes=mes).situacao == "JUSTIFICADO"


def test_mudar_data_recalcula_mes_antigo_e_novo(user_a, filial_a, criar_notificacao):
    n = criar_notificacao(filial_a, user_a, data=datetime.date(2025, 3, 10))
    assert FechamentoMensal.objects.get(filial=filial_a, ano=2025, mes=3).situacao == "COM_NOTIFICACOES"
    n.data_recebimento = datetime.date(2025, 4, 2)
    n.save()
    assert FechamentoMensal.objects.get(filial=filial_a, ano=2025, mes=3).situacao == "PENDENTE"
    assert FechamentoMensal.objects.get(filial=filial_a, ano=2025, mes=4).situacao == "COM_NOTIFICACOES"


def test_justificativa_curta_e_rejeitada(user_a, filial_a, hoje):
    with pytest.raises(ValidationError):
        fech.justificar(user_a, filial_a, hoje.year, hoje.month, "curta")


def test_historico_de_justificativas_e_imutavel(user_a, filial_a, hoje):
    fechamento = fech.justificar(user_a, filial_a, hoje.year, hoje.month, "Primeira justificativa")
    fech.justificar(user_a, filial_a, hoje.year, hoje.month, "Segunda justificativa")
    assert fechamento.justificativas.count() == 2
    j = JustificativaHistorico.objects.first()
    j.texto = "alterada"
    with pytest.raises(ValidationError):
        j.save()
    with pytest.raises(ValidationError):
        j.delete()


# ---------- status vencida e dashboard ----------

def test_vencida_e_derivada_do_prazo(user_a, filial_a, criar_notificacao, hoje):
    atrasada = criar_notificacao(filial_a, user_a, prazo=hoje - datetime.timedelta(days=1), numero="X1")
    respondida = criar_notificacao(filial_a, user_a, prazo=hoje - datetime.timedelta(days=1), status="RESPONDIDA", numero="X2")
    no_prazo = criar_notificacao(filial_a, user_a, prazo=hoje, numero="X3")
    assert atrasada.vencida and not respondida.vencida and not no_prazo.vencida
    assert set(Notificacao.objects.vencidas().values_list("numero_det", flat=True)) == {"X1"}


def test_totais_do_dashboard(user_a, user_b, filial_a, filial_b, criar_notificacao, hoje):
    ano, mes = 2025, 6
    d = datetime.date
    criar_notificacao(filial_a, user_a, data=d(2025, 6, 5), prazo=d(2025, 6, 20), numero="1")
    criar_notificacao(filial_a, user_a, data=d(2025, 6, 9), status="RESPONDIDA", numero="2")
    criar_notificacao(filial_a, user_a, data=d(2025, 2, 1), numero="3")
    criar_notificacao(filial_b, user_b, data=d(2025, 6, 15), numero="4")
    criar_notificacao(filial_b, user_b, data=d(2024, 12, 1), numero="5")  # outro ano

    dados = dash.dados_completos(ano, mes, hoje=d(2025, 6, 18))
    assert dados["resumo"]["total_ano"] == 4
    assert dados["resumo"]["total_mes"] == 3
    assert dados["resumo"]["respondidas"] == 1
    assert dados["resumo"]["vencidas"] == 0  # prazo 20/06 ainda não passou em 18/06
    assert dados["resumo"]["prazo_7_dias"] == 1
    por = {f["filial"].nome: f for f in dados["filiais"]}
    assert (por["Filial A"]["total_ano"], por["Filial A"]["total_mes"]) == (3, 2)
    assert (por["Filial B"]["total_ano"], por["Filial B"]["total_mes"]) == (1, 1)
    assert dados["serie"][5]["total"] == 3 and dados["serie"][1]["total"] == 1
    assert dados["conformidade"]["COM_NOTIFICACOES"] == 2

    # vencida passa a contar depois do prazo (só a notificação "1" tem prazo)
    assert dash.dados_completos(ano, mes, hoje=d(2025, 6, 25))["resumo"]["vencidas"] == 1
